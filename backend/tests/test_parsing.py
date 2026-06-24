"""parse_file 内容层提取测试（**不按大小降级**）。

验证用户指令后的新契约：大小不再决定读不读正文；按类型分派只取内容层；
metadata_only 只在「真提不出内容」时用（需OCR/加密/损坏）；性能保护走超时不走大小。
"""
import time

import app.parsing as parsing
from app.parsing import parse_file


# ── 大小不再是开关 ──────────────────────────────────────────────
def test_no_size_gate_large_text_still_parsed(tmp_path):
    """>20MB 的文本文件仍正常读到正文——证明旧的 20MB 降级开关已彻底移除。"""
    big = tmp_path / "大材料汇编.txt"
    marker = "甲方要求：退台立面与展示区品质。\n"
    with open(big, "w", encoding="utf-8") as fh:
        fh.write(marker)
        fh.write("填充内容。" * 1_500_000)  # ~22MB，远超旧 20MB 阈值
    assert big.stat().st_size > 20 * 1024 * 1024
    r = parse_file(big)
    assert r.status == "ok"            # 没有因为「大」被降级
    assert "退台立面" in r.text         # 真读到了正文
    assert len(r.text) <= parsing.MAX_TEXT_CHARS  # 仍按字符数截断（与大小无关）


def test_small_text_file_ok(tmp_path):
    small = tmp_path / "需求.txt"
    small.write_text("甲方要求：展示区品质与材料质感。", encoding="utf-8")
    r = parse_file(small)
    assert r.status == "ok"
    assert "材料质感" in r.text


def test_empty_text_file_empty(tmp_path):
    empty = tmp_path / "空.txt"
    empty.write_text("   ", encoding="utf-8")
    r = parse_file(empty)
    assert r.status == "empty"


def test_unsupported_ext(tmp_path):
    f = tmp_path / "x.zip"
    f.write_bytes(b"PK\x03\x04stuff")
    assert parse_file(f).status == "unsupported"


# ── PDF 分档：metadata_only 只在真提不出内容时；ok_truncated 如实标截断 ─────
def test_blank_pdf_needs_ocr_metadata_only(tmp_path):
    """有效但无文字层的 PDF（扫描件代表）→ metadata_only + 需OCR，不伪造正文。"""
    import fitz

    f = tmp_path / "扫描件.pdf"
    doc = fitz.open()
    doc.new_page(width=300, height=300)  # 空白页,无文字层
    doc.save(str(f))
    doc.close()
    r = parse_file(f)
    assert r.status == "metadata_only"
    assert r.reason == "needs_ocr"
    assert "需OCR" in r.text and "不伪造" in r.text


def test_corrupt_pdf_unreadable_metadata_only(tmp_path):
    """结构损坏的 PDF → metadata_only + unreadable（如实登记，不伪造、不静默崩）。"""
    f = tmp_path / "损坏.pdf"
    f.write_bytes(b"%PDF-1.4\n" + b"garbage-not-a-real-pdf" * 20)
    r = parse_file(f)
    assert r.status == "metadata_only"
    assert r.reason == "unreadable"
    assert "损坏.pdf" in r.text


def test_pdf_full_text_ok_with_pages(tmp_path):
    """有文字层的多页 PDF 完整读完 → ok，带 total_pages，不截断。
    (fitz 默认字体不嵌 CJK,故用 ASCII 内容验证页/截断逻辑——逻辑与字体无关。)"""
    import fitz

    f = tmp_path / "brief.pdf"
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i+1}: setback facade and showroom quality")
    doc.save(str(f))
    doc.close()
    r = parse_file(f)
    assert r.status == "ok"
    assert r.total_pages == 3
    assert r.truncated_at_page == 0
    assert "setback facade" in r.text


def test_pdf_truncation_honestly_reported(tmp_path, monkeypatch):
    """正文超 MAX_TEXT_CHARS → ok_truncated，如实标停在第N页/共M页（不把截断值当全文）。"""
    import fitz

    monkeypatch.setattr(parsing, "MAX_TEXT_CHARS", 50)  # 调小上限,前几页就触顶
    f = tmp_path / "long.pdf"
    doc = fitz.open()
    for i in range(5):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i+1}: " + "architectural design content " * 5)
    doc.save(str(f))
    doc.close()
    r = parse_file(f)
    assert r.status == "ok_truncated"
    assert 0 < r.truncated_at_page < r.total_pages  # 停在中途某页,非最后一页
    assert r.total_pages == 5
    assert len(r.text) <= 50  # 截断值,不是全文


def test_pdf_truncation_detected_after_nfkc_expansion(tmp_path, monkeypatch):
    """归一化后长度判截断(NFKC 展开/join 换行使真实长度>raw 计数)→ 仍如实标 ok_truncated,
    绝不因 raw 计数没触顶就误报 ok(本次对抗复核坐实的失实点的回归)。"""
    import fitz

    # 上限设到恰好让「join 换行 + 多页」累计超过、但单页 raw 计数边界附近
    monkeypatch.setattr(parsing, "MAX_TEXT_CHARS", 40)
    f = tmp_path / "nfkc.pdf"
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i+1} content block ABCDEFGHIJ")
    doc.save(str(f))
    doc.close()
    r = parse_file(f)
    # 最终归一化文本若超上限,必须 ok_truncated 且 text 不超上限(不把截断值当全文)
    if len(r.text) >= 40 or r.status == "ok_truncated":
        assert r.status == "ok_truncated"
        assert r.truncated_at_page > 0
        assert len(r.text) <= 40


# ── PPTX：抽文本框 + 演讲者备注 ─────────────────────────────────
def test_pptx_extracts_text_and_notes(tmp_path):
    from pptx import Presentation

    f = tmp_path / "汇报.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "退台立面方案要点"
    slide.notes_slide.notes_text_frame.text = "备注：甲方强调檐口比例与材料质感"
    prs.save(str(f))
    r = parse_file(f)
    assert r.status == "ok"
    assert "退台立面方案要点" in r.text   # 文本框
    assert "檐口比例" in r.text           # 演讲者备注（旧逻辑漏抽）


# ── 图片资产：登记为 ok（保持原行为，不受大小影响）─────────────
def test_image_asset_ok_regardless_of_size(tmp_path):
    img = tmp_path / "效果图.png"
    img.write_bytes(b"\x89PNG\r\n" + b"y" * 1000)
    r = parse_file(img)
    assert r.status == "ok"
    assert "图片资产" in r.text


# ── 性能保护：超时（不是大小）兜底 ─────────────────────────────
def test_extraction_timeout_not_size(tmp_path, monkeypatch):
    """异常慢的提取命中超时 → extraction_timeout，证明性能保护走超时而非大小预判。"""
    f = tmp_path / "慢.txt"
    f.write_text("内容", encoding="utf-8")

    def _slow(_p, _ext):
        time.sleep(2)
        return parsing.ParseResult(status="ok", text="不该到这")

    monkeypatch.setattr(parsing, "_dispatch", _slow)
    monkeypatch.setattr(parsing, "PARSE_TIMEOUT_SECONDS", 1)
    r = parse_file(f)
    assert r.status == "extraction_timeout"
    assert r.reason == "timeout"
    assert "未按文件大小放弃" in r.text
