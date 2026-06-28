"""从项目文件抽出图片资产(一等资产):PPT 按 slide / PDF 按 page / Word 嵌入图。

红线对齐 parsing.py：只抽【已存在的嵌入图】，不渲染整页、不 OCR、不伪造。
- Pillow 生成缩略图(长边 ≤ THUMB_MAX)、读尺寸。
- 带上限(MAX_ASSETS_PER_FILE)防超大文件图爆；太小的(图标/项目符号)按 MIN_DIM 跳过。
- 分项失败跳过、整体不抛(不阻塞上传链路)。返回纯数据，落盘/落库由调用方做。
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

THUMB_MAX = 360
MAX_ASSETS_PER_FILE = 80
MIN_DIM = 80          # 长/宽任一 < 此值视作图标/装饰，跳过
CAPTION_MAX = 300


@dataclass
class ExtractedAsset:
    data: bytes
    ext: str
    width: int = 0
    height: int = 0
    page_no: int = 0
    slide_no: int = 0
    shape_index: int = 0
    caption: str = ""


def _dims(data: bytes) -> tuple[int, int]:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as im:
            return int(im.width), int(im.height)
    except Exception:  # noqa: BLE001
        return 0, 0


def make_thumb(data: bytes, max_side: int = THUMB_MAX) -> bytes | None:
    """生成 JPEG 缩略图;失败返回 None(调用方退回用原图)。"""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as im:
            im = im.convert("RGB")
            im.thumbnail((max_side, max_side))
            buf = io.BytesIO()
            im.save(buf, format="JPEG", quality=82)
            return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None


def _norm_ext(blob_ext: str) -> str:
    e = (blob_ext or "").lower().lstrip(".")
    return e if e in ("png", "jpg", "jpeg", "gif", "bmp", "webp", "tiff", "emf", "wmf") else "png"


def _too_small(w: int, h: int) -> bool:
    return bool(w and h) and (w < MIN_DIM or h < MIN_DIM)


def extract_pptx(path: Path) -> list[ExtractedAsset]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[ExtractedAsset] = []
    prs = Presentation(str(path))
    for sidx, slide in enumerate(prs.slides, start=1):
        texts = [
            sh.text_frame.text.strip()
            for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        caption = " ".join(texts)[:CAPTION_MAX]
        for shix, shape in enumerate(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                img = shape.image
                data = img.blob
            except Exception:  # noqa: BLE001
                continue
            w, h = _dims(data)
            if _too_small(w, h):
                continue
            out.append(ExtractedAsset(data=data, ext=_norm_ext(img.ext), width=w, height=h,
                                      slide_no=sidx, shape_index=shix, caption=caption))
            if len(out) >= MAX_ASSETS_PER_FILE:
                return out
    return out


def extract_docx(path: Path) -> list[ExtractedAsset]:
    import docx

    out: list[ExtractedAsset] = []
    d = docx.Document(str(path))
    for rel in d.part.rels.values():
        if "image" not in rel.reltype:
            continue
        try:
            data = rel.target_part.blob
        except Exception:  # noqa: BLE001
            continue
        w, h = _dims(data)
        if _too_small(w, h):
            continue
        out.append(ExtractedAsset(data=data, ext=_norm_ext(Path(str(rel.target_ref)).suffix),
                                  width=w, height=h))
        if len(out) >= MAX_ASSETS_PER_FILE:
            break
    return out


def extract_pdf(path: Path) -> list[ExtractedAsset]:
    import fitz

    out: list[ExtractedAsset] = []
    doc = fitz.open(str(path))
    try:
        for pno in range(doc.page_count):
            page = doc.load_page(pno)
            caption = (page.get_text("text") or "").strip()[:CAPTION_MAX]
            seen: set[int] = set()
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                try:
                    info = doc.extract_image(xref)
                    data = info["image"]
                    ext = _norm_ext(info.get("ext"))
                except Exception:  # noqa: BLE001
                    continue
                w, h = _dims(data)
                if _too_small(w, h):
                    continue
                out.append(ExtractedAsset(data=data, ext=ext, width=w, height=h,
                                          page_no=pno + 1, caption=caption))
                if len(out) >= MAX_ASSETS_PER_FILE:
                    return out
    finally:
        doc.close()
    return out


def extract(path: Path, ext: str) -> list[ExtractedAsset]:
    """按类型抽图;不支持的类型或异常 → 空列表(不抛)。"""
    e = (ext or "").lower()
    try:
        if e == ".pptx":
            return extract_pptx(path)
        if e == ".docx":
            return extract_docx(path)
        if e == ".pdf":
            return extract_pdf(path)
    except Exception:  # noqa: BLE001  抽取整体失败不阻塞上传
        return []
    return []
