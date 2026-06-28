"""会议成果导出（4E）：Word(.docx) 正式交付 + 打印友好 HTML。

- Word 用 python-docx；对外版剔除对内研判，对内版含全部。
- 文件名：项目名_会议名_YYYYMMDD.docx（调用方拼装）。
- 打印 HTML：内联样式 + @media print，前端 window.print() 保存 PDF。
- 不引入后端 PDF 重依赖（WeasyPrint 等）。
"""
from __future__ import annotations

import io
import re
from typing import List

_SEC = ["一、会议背景", "二、关键结论", "三、甲方诉求", "四、风险与分歧", "五、下一步行动"]


def safe_filename(*parts: str) -> str:
    name = "_".join(p.strip() for p in parts if p and p.strip())
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    return (name or "会议纪要")[:120]


def build_pptx(data: dict) -> bytes:
    """把 PPT 大纲结构化数据(skill_structured.normalize_ppt 的输出)渲染成真 .pptx。

    版式:封面页 + 每页(页标题 + 核心观点 + 要点);讲稿提示/来源进幻灯片备注。
    visualSuggestion 渲染为右侧占位框——这是【图片资产层】的接入点:以后某页带 image_ref
    时,把占位换成 slide.shapes.add_picture(asset_path) 即可,版式无需重排。
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W: int = prs.slide_width

    def textbox(slide, left, top, width, height):
        tf = slide.shapes.add_textbox(left, top, width, height).text_frame
        tf.word_wrap = True
        return tf

    def line(tf, text: str, *, size: int, bold: bool = False, color=None, first: bool = False):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        if color is not None:
            p.font.color.rgb = RGBColor(*color)
        return p

    GREY = (0x88, 0x88, 0x88)
    LGREY = (0xA8, 0xA8, 0xA8)
    TERRA = (0xB0, 0x55, 0x2A)

    # ── 封面 ──
    cover = prs.slides.add_slide(blank)
    tf = textbox(cover, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(2.4))
    line(tf, data.get("title") or "汇报", size=40, bold=True, first=True)
    if data.get("subtitle"):
        line(tf, data["subtitle"], size=20, color=GREY)
    meta = []
    if data.get("audience"):
        meta.append("汇报对象：" + data["audience"])
    meta.append("共 %d 页" % len(data.get("slides") or []))
    line(tf, "　".join(meta), size=14, color=LGREY)
    if data.get("narrative"):
        nf = textbox(cover, Inches(0.9), Inches(5.2), W - Inches(1.8), Inches(1.6))
        line(nf, data["narrative"], size=13, color=GREY, first=True)

    # ── 内容页 ──
    for s in data.get("slides") or []:
        slide = prs.slides.add_slide(blank)
        head = textbox(slide, Inches(0.6), Inches(0.4), W - Inches(1.2), Inches(0.9))
        line(head, f"{s.get('no', '')}. {s.get('title', '')}".strip(". "), size=26, bold=True, first=True)

        body_top = Inches(1.5)
        if s.get("keyMessage"):
            kf = textbox(slide, Inches(0.6), body_top, W - Inches(1.2), Inches(0.8))
            line(kf, "▎ " + s["keyMessage"], size=16, bold=True, color=TERRA, first=True)
            body_top = Inches(2.4)

        bf = textbox(slide, Inches(0.6), body_top, Emu(int(W * 0.56)), Inches(4.0))
        bullets = s.get("bullets") or []
        for i, b in enumerate(bullets):
            line(bf, "• " + str(b), size=14, first=(i == 0))

        # 视觉建议 → 右侧占位(图片层接入点)
        if s.get("visualSuggestion"):
            vf = textbox(slide, Emu(int(W * 0.63)), Inches(1.5), Emu(int(W * 0.33)), Inches(4.5))
            line(vf, "🖼 建议图：", size=12, bold=True, color=LGREY, first=True)
            line(vf, str(s["visualSuggestion"]), size=12, color=LGREY)

        notes = []
        if s.get("speakerNotes"):
            notes.append("讲稿：" + s["speakerNotes"])
        if s.get("sourceRefs"):
            notes.append("来源：" + "、".join(s["sourceRefs"]))
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_docx(title: str, m: dict, *, internal: bool) -> bytes:
    """生成 .docx 字节。internal=False→对外版(不含对内研判)；True→对内版(含全部)。"""
    from docx import Document

    doc = Document()
    doc.add_heading(f"会议纪要 · {title}", level=0)
    variant = "对内研判版（不对外）" if internal else "对外纪要版"
    doc.add_paragraph(variant)

    doc.add_heading(_SEC[0], level=1)
    for x in m.get("summary", []):
        doc.add_paragraph(str(x), style="List Bullet")

    doc.add_heading(_SEC[1], level=1)
    for x in m.get("core_items", []):
        doc.add_paragraph(str(x), style="List Bullet")

    doc.add_heading(_SEC[2], level=1)
    demands = m.get("demand_internal", []) if internal else m.get("demand_external", [])
    for d in demands:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(d.get("statement", ""))
        q = d.get("quote", "")
        if q:
            run = p.add_run(f"（原话「{q}」 {d.get('time','')}）")
            run.italic = True

    doc.add_heading(_SEC[3], level=1)
    for x in m.get("decisions", []):
        doc.add_paragraph(str(x), style="List Bullet")

    doc.add_heading(_SEC[4], level=1)
    for t in m.get("todos", []):
        owner = f" @{t.get('owner','')}" if t.get("owner") else ""
        due = f"（{t.get('due','')}）" if t.get("due") else ""
        doc.add_paragraph(f"☐ {t.get('text','')}{owner}{due}", style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _esc(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def build_print_html(title: str, m: dict, *, internal: bool = False) -> str:
    """打印友好 HTML（@media print）。前端用 window.print() 保存 PDF。"""
    demands = m.get("demand_internal", []) if internal else m.get("demand_external", [])
    variant = "对内研判版（不对外）" if internal else "对外纪要版"

    def ul(items):
        return "<ul>" + "".join(f"<li>{_esc(x)}</li>" for x in items) + "</ul>" if items else "<p class='muted'>—</p>"

    demand_html = "<ul>" + "".join(
        f"<li>{_esc(d.get('statement',''))}"
        + (f" <span class='q'>（原话「{_esc(d.get('quote',''))}」 {_esc(d.get('time',''))}）</span>" if d.get("quote") else "")
        + "</li>"
        for d in demands
    ) + "</ul>" if demands else "<p class='muted'>—</p>"

    todos_html = "<ul>" + "".join(
        f"<li>☐ {_esc(t.get('text',''))}"
        + (f" @{_esc(t.get('owner',''))}" if t.get("owner") else "")
        + (f"（{_esc(t.get('due',''))}）" if t.get("due") else "")
        + "</li>"
        for t in m.get("todos", [])
    ) + "</ul>" if m.get("todos") else "<p class='muted'>—</p>"

    return f"""<!doctype html><html lang="zh"><head><meta charset="utf-8">
<title>会议纪要 · {_esc(title)}</title>
<style>
  body {{ font-family: "Microsoft YaHei", "PingFang SC", sans-serif; color:#1d1b16; max-width:780px; margin:32px auto; padding:0 16px; line-height:1.7; }}
  h1 {{ font-size:22px; border-bottom:2px solid #b07a35; padding-bottom:8px; }}
  h2 {{ font-size:16px; margin-top:22px; color:#7a5a2a; }}
  .variant {{ color:#888; font-size:13px; }}
  .q {{ color:#888; font-size:12px; }}
  .muted {{ color:#aaa; }}
  ul {{ padding-left:20px; }}
  .toolbar {{ position:fixed; top:12px; right:12px; }}
  button {{ padding:8px 14px; border:0; border-radius:8px; background:#1d1b16; color:#fff; cursor:pointer; }}
  @media print {{ .toolbar {{ display:none; }} body {{ margin:0; }} }}
</style></head><body>
<div class="toolbar"><button onclick="window.print()">打印 / 保存 PDF</button></div>
<h1>会议纪要 · {_esc(title)}</h1>
<p class="variant">{variant}</p>
<h2>{_SEC[0]}</h2>{ul(m.get("summary", []))}
<h2>{_SEC[1]}</h2>{ul(m.get("core_items", []))}
<h2>{_SEC[2]}</h2>{demand_html}
<h2>{_SEC[3]}</h2>{ul(m.get("decisions", []))}
<h2>{_SEC[4]}</h2>{todos_html}
</body></html>"""
